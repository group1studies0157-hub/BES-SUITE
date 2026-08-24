"""
Parses bore-log workbooks (.xls or .xlsx), one sheet per borehole.
See table_common.py for the shared row/column parsing logic - this
module is only responsible for getting rows of cell values out of the
xls/xlsx file.
"""
from __future__ import annotations

from .model import BoreholeRecord
from .table_common import parse_grid


class WorkbookReader:
    """Small abstraction so we can support both xlrd (.xls) and
    openpyxl (.xlsx) behind one interface."""

    def __init__(self, path: str):
        self.path = path
        if path.lower().endswith('.xls'):
            import xlrd
            self._wb = xlrd.open_workbook(path)
            self._kind = 'xlrd'
        else:
            import openpyxl
            self._wb = openpyxl.load_workbook(path, data_only=True)
            self._kind = 'openpyxl'

    def sheet_names(self):
        if self._kind == 'xlrd':
            return self._wb.sheet_names()
        return self._wb.sheetnames

    def rows(self, sheet_name: str):
        """Return rows as a list of lists of raw cell values."""
        if self._kind == 'xlrd':
            sheet = self._wb.sheet_by_name(sheet_name)
            return [[sheet.cell_value(r, c) for c in range(sheet.ncols)]
                    for r in range(sheet.nrows)]
        ws = self._wb[sheet_name]
        return [[('' if v is None else v) for v in row]
                for row in ws.iter_rows(values_only=True)]


def parse_workbook(path: str) -> tuple[list[BoreholeRecord], list[str]]:
    """Parse every sheet of the workbook as one borehole each.
    Returns (boreholes, warnings)."""
    warnings: list[str] = []
    reader = WorkbookReader(path)
    boreholes = []
    for name in reader.sheet_names():
        rows = reader.rows(name)
        try:
            bh = parse_grid(name, rows, warnings)
        except Exception as e:
            warnings.append(f"[sheet {name}] failed to parse: {e}")
            bh = None
        if bh is not None:
            boreholes.append(bh)
    return boreholes, warnings
